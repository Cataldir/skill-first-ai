"""Build an editable PPTX deck for the AI Tech Circle talk.

Run from the repo root:

    python presentation/build_pptx.py

Output: ``presentation/skill-first-ai.pptx`` with native, editable text
boxes (no rasterized slides). Designed with Microsoft visual identity:
Segoe UI typography and the Microsoft web color palette.

The official Microsoft four-square logo is **not** embedded — using a
text wordmark for redistribution safety. Drop the official PNG over the
top-left wordmark on the cover when presenting from a Microsoft device.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

# ----------------------------------------------------------------------
# Microsoft palette + typography
# ----------------------------------------------------------------------

MS_BLUE = RGBColor(0x00, 0x78, 0xD4)
MS_DARK = RGBColor(0x00, 0x20, 0x50)
MS_DEEP = RGBColor(0x00, 0x10, 0x30)
MS_ACCENT = RGBColor(0x50, 0xE6, 0xFF)
MS_GRAY = RGBColor(0xF3, 0xF2, 0xF1)
MS_LINE = RGBColor(0xD2, 0xD0, 0xCE)
MS_INK = RGBColor(0x20, 0x1F, 0x1E)
MS_MUTED = RGBColor(0x60, 0x5E, 0x5C)
MS_WARN = RGBColor(0xD8, 0x3B, 0x01)
MS_GOOD = RGBColor(0x10, 0x7C, 0x10)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT_HEAD = "Segoe UI Semibold"
FONT_BODY = "Segoe UI"
FONT_MONO = "Cascadia Mono"

EVENT_LABEL = "AI Tech Circle · 28 mai 2026"
DECK_LABEL = "Skill-first AI"
SPEAKER = "Ricardo Cataldi · Microsoft GBB"
REPO_URL = "github.com/Cataldir/skill-first-ai"

# 16:9 widescreen
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

# Content rectangle
MARGIN_X = Inches(0.7)
CONTENT_W = SLIDE_W - 2 * MARGIN_X
TITLE_TOP = Inches(0.95)
TITLE_H = Inches(1.05)
BODY_TOP = Inches(2.15)
BODY_H = Inches(4.55)


# ----------------------------------------------------------------------
# Drawing helpers
# ----------------------------------------------------------------------


def _solid_rect(slide, left, top, width, height, color):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def _paint_background(slide, color):
    """Add a full-bleed background rectangle behind everything else."""
    shp = _solid_rect(slide, 0, 0, SLIDE_W, SLIDE_H, color)
    spTree = shp._element.getparent()
    spTree.remove(shp._element)
    spTree.insert(2, shp._element)
    return shp


def _textbox(slide, left, top, width, height):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    return box


def _set_run(run, *, text, font=FONT_BODY, size=18, color=MS_INK,
             bold=False, italic=False):
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def _para_run(p, *, text, font=FONT_BODY, size=18, color=MS_INK,
              bold=False, italic=False):
    run = p.add_run()
    _set_run(run, text=text, font=font, size=size, color=color,
             bold=bold, italic=italic)
    return run


def add_text(slide, left, top, width, height, text, *, font=FONT_BODY,
             size=18, color=MS_INK, bold=False, italic=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = _textbox(slide, left, top, width, height)
    tf = box.text_frame
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    _para_run(p, text=text, font=font, size=size, color=color,
              bold=bold, italic=italic)
    return box


def add_paragraphs(slide, left, top, width, height, lines, *,
                   default_font=FONT_BODY, default_size=18,
                   default_color=MS_INK, align=PP_ALIGN.LEFT,
                   anchor=MSO_ANCHOR.TOP, line_spacing=1.25,
                   space_before_pt=0):
    """Add a multi-paragraph text box.

    Each entry in ``lines`` is either a plain ``str`` or a ``(text, opts)``
    tuple. Opts can carry ``bold``, ``italic``, ``color``, ``size``,
    ``font``, ``align``, ``bullet``, and ``space_before``.
    """
    box = _textbox(slide, left, top, width, height)
    tf = box.text_frame
    tf.vertical_anchor = anchor

    for i, entry in enumerate(lines):
        if isinstance(entry, tuple):
            text, opts = entry
        else:
            text, opts = entry, {}

        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = opts.get("align", align)
        p.line_spacing = line_spacing
        if "space_before" in opts:
            p.space_before = Pt(opts["space_before"])
        elif i > 0 and space_before_pt:
            p.space_before = Pt(space_before_pt)

        if opts.get("bullet"):
            text = f"•  {text}"

        _para_run(
            p,
            text=text,
            font=opts.get("font", default_font),
            size=opts.get("size", default_size),
            color=opts.get("color", default_color),
            bold=opts.get("bold", False),
            italic=opts.get("italic", False),
        )
    return box


def add_accent_bar(slide, top=Inches(2.05), color=MS_BLUE,
                   width=Inches(1.0), height=Inches(0.07)):
    return _solid_rect(slide, MARGIN_X, top, width, height, color)


def add_chip(slide, left, top, label, *, fill=MS_BLUE, color=WHITE,
             padding_x=Inches(0.18), height=Inches(0.34), size=12):
    """Add a pill-shaped label chip."""
    approx_w = Pt(size).pt * 0.135 * len(label) * 12700 + 2 * padding_x
    width = max(Inches(1.1), int(approx_w))
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top,
                                 width, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _para_run(p, text=label, font=FONT_HEAD, size=size, color=color,
              bold=True)
    return shp, width


def add_header_footer(slide, page_num, total, *, dark=False):
    if dark:
        ms_color = WHITE
        sub = MS_GRAY
        muted = MS_GRAY
    else:
        ms_color = MS_DARK
        sub = MS_MUTED
        muted = MS_MUTED

    add_text(slide, MARGIN_X, Inches(0.32), Inches(2.5), Inches(0.3),
             "Microsoft", font=FONT_HEAD, size=11,
             color=ms_color, bold=True)
    add_text(slide, Inches(2.7), Inches(0.32), Inches(7.0), Inches(0.3),
             DECK_LABEL, font=FONT_BODY, size=11, color=sub)
    add_text(slide, Inches(8.3), Inches(0.32), SLIDE_W - Inches(9.0),
             Inches(0.3), EVENT_LABEL, font=FONT_BODY, size=11,
             color=sub, align=PP_ALIGN.RIGHT)

    add_text(slide, MARGIN_X, Inches(7.06), Inches(8), Inches(0.3),
             f"{SPEAKER}  ·  {REPO_URL}",
             font=FONT_BODY, size=10, color=muted)
    add_text(slide, Inches(12.2), Inches(7.06), Inches(0.6), Inches(0.3),
             f"{page_num} / {total}", font=FONT_HEAD, size=10,
             color=muted, align=PP_ALIGN.RIGHT)


def add_title(slide, text, *, color=MS_DARK):
    add_text(slide, MARGIN_X, TITLE_TOP, CONTENT_W, TITLE_H, text,
             font=FONT_HEAD, size=36, color=color,
             anchor=MSO_ANCHOR.MIDDLE)
    add_accent_bar(slide)


# ----------------------------------------------------------------------
# Slide builders
# ----------------------------------------------------------------------


def build_cover(slide, total):
    _paint_background(slide, MS_DEEP)

    # Decorative accent stripe along bottom
    _solid_rect(slide, 0, Inches(7.1), SLIDE_W, Inches(0.07), MS_BLUE)
    _solid_rect(slide, 0, Inches(7.17), SLIDE_W, Inches(0.03), MS_ACCENT)

    # Microsoft wordmark — top-left
    add_text(slide, MARGIN_X, Inches(0.5), Inches(3), Inches(0.4),
             "Microsoft", font=FONT_HEAD, size=14,
             color=WHITE, bold=True)

    # Event chip — top-right
    add_chip(slide, Inches(10.5), Inches(0.5), "AI Tech Circle",
             fill=MS_BLUE, color=WHITE, size=12)

    # Main title block
    add_text(slide, MARGIN_X, Inches(2.1), Inches(12), Inches(1.4),
             "Skill-first AI",
             font=FONT_HEAD, size=72, color=WHITE)

    add_text(slide, MARGIN_X, Inches(3.6), Inches(12), Inches(0.9),
             "Menos agentes. Mais capacidades inteligentes.",
             font=FONT_BODY, size=30, color=MS_ACCENT, italic=True)

    # Accent rule
    _solid_rect(slide, MARGIN_X, Inches(4.7), Inches(1.5), Inches(0.06),
                MS_ACCENT)

    # Speaker block
    add_paragraphs(
        slide,
        MARGIN_X, Inches(4.95), Inches(8), Inches(1.8),
        [
            ("Ricardo Cataldi",
             {"bold": True, "size": 24, "color": WHITE}),
            ("GBB Senior Cloud Solution Architect · Microsoft",
             {"size": 16, "color": MS_GRAY, "space_before": 4}),
            ("Conselheiro acadêmico · FIAP / IBM",
             {"size": 16, "color": MS_GRAY, "space_before": 2}),
        ],
        line_spacing=1.15,
    )

    # Event + date bottom-right
    add_paragraphs(
        slide,
        Inches(8.0), Inches(5.6), Inches(4.6), Inches(1.2),
        [
            ("AI Tech Circle",
             {"bold": True, "size": 18, "color": WHITE,
              "align": PP_ALIGN.RIGHT}),
            ("28 de maio de 2026 · São Paulo",
             {"size": 14, "color": MS_GRAY,
              "align": PP_ALIGN.RIGHT, "space_before": 4}),
            (REPO_URL,
             {"size": 13, "color": MS_ACCENT, "italic": True,
              "align": PP_ALIGN.RIGHT, "space_before": 4}),
        ],
        line_spacing=1.15,
    )


# ----------------------------------------------------------------------
# Content layout primitives
# ----------------------------------------------------------------------


def light_slide(slide, total, page):
    _paint_background(slide, WHITE)
    add_header_footer(slide, page, total, dark=False)


def dark_slide(slide, total, page):
    _paint_background(slide, MS_DEEP)
    # Accent stripe
    _solid_rect(slide, MARGIN_X, Inches(2.05), Inches(1.0),
                Inches(0.07), MS_ACCENT)
    add_header_footer(slide, page, total, dark=True)


def question_slide(slide, total, page):
    _paint_background(slide, MS_GRAY)
    add_header_footer(slide, page, total, dark=False)


# ----------------------------------------------------------------------
# Table helper
# ----------------------------------------------------------------------


def add_table(slide, left, top, width, height, headers, rows, *,
              header_fill=MS_DARK, header_color=WHITE,
              row_alt=RGBColor(0xFA, 0xFA, 0xFA),
              size=14, head_size=14):
    cols = len(headers)
    rows_total = len(rows) + 1
    table_shape = slide.shapes.add_table(rows_total, cols, left, top,
                                         width, height)
    table = table_shape.table

    for ci, h in enumerate(headers):
        cell = table.cell(0, ci)
        cell.fill.solid()
        cell.fill.fore_color.rgb = header_fill
        cell.margin_left = Inches(0.12)
        cell.margin_right = Inches(0.12)
        cell.margin_top = Inches(0.06)
        cell.margin_bottom = Inches(0.06)
        cell.text_frame.clear()
        p = cell.text_frame.paragraphs[0]
        _para_run(p, text=h, font=FONT_HEAD, size=head_size,
                  color=header_color, bold=True)

    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            cell = table.cell(ri, ci)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if ri % 2 else row_alt
            cell.margin_left = Inches(0.12)
            cell.margin_right = Inches(0.12)
            cell.margin_top = Inches(0.05)
            cell.margin_bottom = Inches(0.05)
            cell.text_frame.clear()
            p = cell.text_frame.paragraphs[0]
            _para_run(p, text=val, font=FONT_BODY, size=size,
                      color=MS_INK)
    return table


# ----------------------------------------------------------------------
# Individual slides
# ----------------------------------------------------------------------


def slide_pergunta_honesta(slide):
    add_title(slide, "Antes de começar — uma pergunta honesta")
    add_paragraphs(
        slide, MARGIN_X, BODY_TOP, CONTENT_W, BODY_H,
        [
            ("Quantos agentes a sua empresa já criou nos últimos 12 meses?",
             {"size": 28, "color": MS_DARK, "italic": True,
              "space_before": 0}),
            ("Conte os pilotos. Conte os PoCs. Conte os “Copilot for X” "
             "que TI fez para uma diretoria. Conte os bots de RH, "
             "financeiro, jurídico, atendimento.",
             {"size": 20, "space_before": 24}),
            ("Não é trick question. É o ponto de partida.",
             {"size": 20, "bold": True, "color": MS_INK,
              "space_before": 18}),
            ("Se a resposta começa com “uns sete, oito…”, a próxima "
             "pergunta importa mais do que parece:",
             {"size": 20, "space_before": 18}),
            ("quantos deles fazem as mesmas seis coisas em prompts "
             "ligeiramente diferentes?",
             {"size": 22, "bold": True, "color": MS_BLUE,
              "space_before": 12}),
        ],
        line_spacing=1.3,
    )


def slide_armadilha(slide):
    add_title(slide, "A armadilha do “um agente por área”")
    add_chip(slide, MARGIN_X, Inches(2.15),
             "Padrão observado em quase toda enterprise hoje",
             fill=MS_WARN, color=WHITE, size=11)
    headers = ["Área", "Agente", "O que ele faz por dentro"]
    rows = [
        ["RH", "Agente de RH",
         "consulta política · valida permissão · busca documento · abre chamado"],
        ["Financeiro", "Agente Financeiro",
         "consulta política · valida permissão · busca documento · abre chamado"],
        ["Jurídico", "Agente Jurídico",
         "consulta política · valida permissão · busca documento · abre chamado"],
        ["TI", "Agente de Acessos",
         "consulta política · valida permissão · busca documento · abre chamado"],
        ["Atendimento", "Agente CS",
         "consulta política · valida permissão · busca documento · abre chamado"],
    ]
    add_table(slide, MARGIN_X, Inches(2.7), CONTENT_W, Inches(3.4),
              headers, rows, size=13, head_size=13)
    add_text(slide, MARGIN_X, Inches(6.25), CONTENT_W, Inches(0.6),
             "Cinco agentes. Vinte cópias das mesmas quatro capacidades. "
             "Cinco prompts ligeiramente diferentes. Zero auditoria compartilhada.",
             font=FONT_BODY, size=16, color=MS_DARK, italic=True)


def slide_vocabulario(slide):
    add_title(slide, "Vocabulário — antes da gente concordar, separar")
    headers = ["Termo", "O que é", "O que NÃO é"]
    rows = [
        ["Automação tradicional", "fluxo determinístico, regras fixas",
         "inteligente"],
        ["Skill (capacidade)", "função tipada, versionada, governada",
         "conversa"],
        ["Workflow", "orquestração explícita de skills", "adivinhação"],
        ["Agente",
         "composição autônoma de skills + memória", "bot de FAQ"],
    ]
    add_table(slide, MARGIN_X, Inches(2.25), CONTENT_W, Inches(2.8),
              headers, rows, size=16, head_size=16)
    add_text(slide, MARGIN_X, Inches(5.4), CONTENT_W, Inches(1.6),
             "Quando o time fala “agente”, costuma estar misturando os "
             "quatro. Cada um exige um nível diferente de governança.",
             font=FONT_BODY, size=18, color=MS_DARK, italic=True)


def slide_decomposicao(slide):
    add_title(slide, "Decompondo um “agente financeiro” genérico")
    add_text(slide, MARGIN_X, Inches(2.2), CONTENT_W, Inches(0.7),
             "Antes de pedir “crie um agente financeiro”, responda:",
             font=FONT_BODY, size=18, color=MS_INK)
    add_text(slide, MARGIN_X, Inches(2.7), CONTENT_W, Inches(0.7),
             "quais capacidades já deveriam existir?",
             font=FONT_HEAD, size=22, color=MS_BLUE, italic=True)
    add_paragraphs(
        slide, MARGIN_X, Inches(3.6), Inches(7), Inches(3.3),
        [
            ("1.  policy_lookup — consultar política",
             {"size": 19, "font": FONT_BODY}),
            ("2.  permission_check — validar permissão",
             {"size": 19, "font": FONT_BODY, "space_before": 6}),
            ("3.  document_search — buscar documentos",
             {"size": 19, "font": FONT_BODY, "space_before": 6}),
            ("4.  justification_draft — gerar justificativa",
             {"size": 19, "font": FONT_BODY, "space_before": 6}),
            ("5.  ticket_open — abrir chamado",
             {"size": 19, "font": FONT_BODY, "space_before": 6}),
            ("6.  evidence_log — registrar evidência",
             {"size": 19, "font": FONT_BODY, "space_before": 6}),
        ],
        line_spacing=1.2,
    )
    add_paragraphs(
        slide, Inches(8.0), Inches(3.6), Inches(4.6), Inches(3.3),
        [
            ("Nenhuma é específica de finanças.",
             {"size": 18, "bold": True, "color": MS_DARK}),
            ("RH precisa das seis.",
             {"size": 16, "color": MS_INK, "space_before": 10}),
            ("TI precisa das seis.",
             {"size": 16, "color": MS_INK, "space_before": 4}),
            ("Jurídico precisa das seis.",
             {"size": 16, "color": MS_INK, "space_before": 4}),
            ("Atendimento precisa das seis.",
             {"size": 16, "color": MS_INK, "space_before": 4}),
        ],
        line_spacing=1.2,
    )


def slide_demo(slide, *, num, title, bullets, return_cue):
    add_chip(slide, MARGIN_X, Inches(2.25), f"DEMO {num}",
             fill=MS_BLUE, color=WHITE, size=13)
    add_chip(slide, Inches(1.95), Inches(2.25), "demo segmentada",
             fill=MS_DARK, color=MS_ACCENT, size=11)
    add_text(slide, MARGIN_X, Inches(2.75), CONTENT_W, Inches(0.9),
             title, font=FONT_HEAD, size=36, color=WHITE)
    add_paragraphs(
        slide, MARGIN_X, Inches(3.85), Inches(8.5), Inches(2.4),
        bullets, default_color=MS_GRAY, default_size=17,
        line_spacing=1.3, space_before_pt=6,
    )

    # Return-to-deck cue box on the right
    cue_box = _solid_rect(
        slide, Inches(9.2), Inches(3.85), Inches(3.6), Inches(2.4),
        MS_DARK,
    )
    cue_box.line.color.rgb = MS_ACCENT
    cue_box.line.width = Pt(1.5)

    add_paragraphs(
        slide,
        Inches(9.4), Inches(4.0), Inches(3.2), Inches(2.1),
        [
            ("⟵ voltar para o deck",
             {"bold": True, "size": 13, "color": MS_ACCENT}),
            (return_cue,
             {"size": 14, "color": WHITE, "italic": True,
              "space_before": 10}),
        ],
        line_spacing=1.3,
    )


def slide_demo1(slide):
    slide_demo(
        slide,
        num=1,
        title="Foundry portal — finance-orchestrator",
        bullets=[
            "Um prompt agent chamado finance-orchestrator.",
            "Seis tools MCP, expostas via APIM AI Gateway, vindas do registry.",
            "Um turno: “quero reembolso de USD 1800 para visitar um cliente novo”.",
            "Tool calls aparecem em sequência: policy_lookup → permission_check → "
            "document_search → justification_draft → ticket_open → evidence_log.",
        ],
        return_cue=(
            "O que faz aquelas seis ferramentas serem reutilizáveis?"
        ),
    )


def slide_contrato(slide):
    add_title(slide, "O contrato — o que faz uma capability ser reutilizável")
    add_text(slide, MARGIN_X, Inches(2.2), CONTENT_W, Inches(0.5),
             "Quatro trabalhos. Sempre os mesmos.",
             font=FONT_BODY, size=18, color=MS_MUTED, italic=True)

    code = (
        "class BaseSkill(abc.ABC, Generic[InputT, OutputT]):\n"
        "    manifest: ClassVar[SkillManifest]            # identidade\n"
        "    InputSchema: ClassVar[type[BaseModel]]       # tradução\n"
        "    OutputSchema: ClassVar[type[BaseModel]]      # tradução\n"
        "\n"
        "    def run(self, payload, *, correlation_id=None) -> SkillResult:\n"
        "        ...   # falha categorizada + trace de evidência"
    )
    code_box = _solid_rect(slide, MARGIN_X, Inches(2.85),
                           CONTENT_W, Inches(2.0), MS_DEEP)
    code_box.line.fill.background()
    add_text(slide, Inches(0.95), Inches(2.95), CONTENT_W - Inches(0.5),
             Inches(1.8), code, font=FONT_MONO, size=14, color=MS_ACCENT)

    add_paragraphs(
        slide, MARGIN_X, Inches(5.05), CONTENT_W, Inches(1.85),
        [
            ("Identidade — nome estável + versão semântica que dá pra fixar.",
             {"size": 16, "bullet": True}),
            ("Tradução — Pydantic converte o que o agente mandou em modelo de domínio.",
             {"size": 16, "bullet": True, "space_before": 6}),
            ("Falha categorizada — OK · INVALID_INPUT · REFUSED · TEMPORARY_FAILURE.",
             {"size": 16, "bullet": True, "space_before": 6}),
            ("Evidência — SkillTrace com correlation id, latência, categoria.",
             {"size": 16, "bullet": True, "space_before": 6}),
        ],
        line_spacing=1.3,
    )


def slide_envelope(slide):
    add_title(slide, "O resultado é monótono — e isso é uma qualidade")

    code = (
        "class SkillResult(BaseModel, Generic[OutputT]):\n"
        "    category: ResultCategory\n"
        "    output: OutputT | None\n"
        "    message: str\n"
        "    trace: SkillTrace"
    )
    code_box = _solid_rect(slide, MARGIN_X, Inches(2.25),
                           CONTENT_W, Inches(1.85), MS_DEEP)
    code_box.line.fill.background()
    add_text(slide, Inches(0.95), Inches(2.4), CONTENT_W - Inches(0.5),
             Inches(1.7), code, font=FONT_MONO, size=15, color=MS_ACCENT)

    add_text(slide, MARGIN_X, Inches(4.4), CONTENT_W, Inches(0.7),
             "Todo consumidor — agente Foundry, MCP server, FastAPI "
             "interno, import direto em Python — recebe o mesmo envelope.",
             font=FONT_BODY, size=18, color=MS_INK)

    quote_box = _solid_rect(slide, MARGIN_X, Inches(5.4),
                            CONTENT_W, Inches(1.3), MS_GRAY)
    quote_box.line.fill.background()
    _solid_rect(slide, MARGIN_X, Inches(5.4), Inches(0.07),
                Inches(1.3), MS_BLUE)
    add_text(slide, Inches(0.95), Inches(5.55), CONTENT_W - Inches(0.5),
             Inches(1.0),
             "Se a sua skill precisa de um “modo especial” para um agente "
             "específico, ela ainda não é uma skill. É um pedaço do "
             "agente disfarçado.",
             font=FONT_BODY, size=18, color=MS_DARK, italic=True)


def slide_demo2(slide):
    slide_demo(
        slide,
        num=2,
        title="Repositório — onde mora o contrato",
        bullets=[
            "src/skill_first_ai/contracts/skill.py — BaseSkill, SkillManifest, ResultCategory.",
            "src/skill_first_ai/skills/policy_lookup.py — uma skill real, ~70 linhas.",
            "src/skill_first_ai/registry/registry.py — registry + default_registry().",
            "src/skill_first_ai/foundry/mcp_server.py — vira MCP server em um único arquivo.",
        ],
        return_cue=(
            "Como a gente impede que essas skills virem bagunça em seis meses?"
        ),
    )


def slide_governanca(slide):
    add_title(slide, "Governança — testar, versionar, governar")
    items = [
        ("testar",
         "Schema golden",
         "qualquer mudança no input é detectada como drift e exige bump de versão."),
        ("versionar",
         "Semver no manifesto",
         "agentes fixam versão. Skill nova nasce em 1.0.0; breaking change vira 2.0.0."),
        ("governar",
         "Manifesto obrigatório",
         "sem owner, descrição, side_effects ou pii declarados, a skill nem entra no registry."),
    ]
    y = Inches(2.25)
    for chip_label, head, body in items:
        chip, chip_w = add_chip(slide, MARGIN_X, y, chip_label,
                                fill=MS_DARK, color=MS_ACCENT, size=11)
        add_text(slide, MARGIN_X + chip_w + Inches(0.15), y - Inches(0.04),
                 Inches(4), Inches(0.4),
                 head, font=FONT_HEAD, size=18, color=MS_DARK, bold=True)
        add_text(slide, MARGIN_X + Inches(0.05), y + Inches(0.45),
                 CONTENT_W, Inches(0.6),
                 body, font=FONT_BODY, size=15, color=MS_INK)
        y = y + Inches(1.15)

    quote_box = _solid_rect(slide, MARGIN_X, Inches(6.05),
                            CONTENT_W, Inches(0.85), MS_GRAY)
    quote_box.line.fill.background()
    _solid_rect(slide, MARGIN_X, Inches(6.05), Inches(0.07),
                Inches(0.85), MS_BLUE)
    add_text(slide, Inches(0.95), Inches(6.15), CONTENT_W - Inches(0.5),
             Inches(0.7),
             "A barreira não é técnica. É social. O contrato escrito "
             "força o time a decidir antes de publicar.",
             font=FONT_BODY, size=17, color=MS_DARK, italic=True)


def slide_demo3(slide):
    slide_demo(
        slide,
        num=3,
        title="Governança rodando — pytest",
        bullets=[
            "uv run pytest tests/test_governance.py -v",
            "test_skill_has_complete_manifest — owner, descrição, versão.",
            "test_skill_declares_schemas — extra='forbid' no input.",
            "test_side_effecting_skills_require_pii_review — side-effect implica PII review.",
            "test_skill_input_schema_matches_golden — snapshot do JSON schema.",
        ],
        return_cue=(
            "Por que isso muda o que a gente discute em comitê?"
        ),
    )


def slide_reuso(slide):
    add_title(slide, "Reuso real — o teste do segundo agente")
    add_text(slide, MARGIN_X, Inches(2.2), CONTENT_W, Inches(0.5),
             "Mesmo registry. Outro system_prompt. Agente de RH pronto.",
             font=FONT_BODY, size=18, color=MS_INK)

    code = (
        "name: hr-orchestrator\n"
        "tools:\n"
        "  - policy_lookup        # mesma skill, manifesto idêntico\n"
        "  - permission_check     # mesma skill, manifesto idêntico\n"
        "  - document_search      # mesma skill, manifesto idêntico\n"
        "  - justification_draft  # mesma skill, manifesto idêntico\n"
        "  - ticket_open          # mesma skill, requires_human_approval: true\n"
        "  - evidence_log         # mesma skill, append-only"
    )
    code_box = _solid_rect(slide, MARGIN_X, Inches(2.85),
                           CONTENT_W, Inches(2.6), MS_DEEP)
    code_box.line.fill.background()
    add_text(slide, Inches(0.95), Inches(2.95), CONTENT_W - Inches(0.5),
             Inches(2.4), code, font=FONT_MONO, size=15, color=MS_ACCENT)

    quote_box = _solid_rect(slide, MARGIN_X, Inches(5.7),
                            CONTENT_W, Inches(1.1), MS_GRAY)
    quote_box.line.fill.background()
    _solid_rect(slide, MARGIN_X, Inches(5.7), Inches(0.07),
                Inches(1.1), MS_BLUE)
    add_text(slide, Inches(0.95), Inches(5.85), CONTENT_W - Inches(0.5),
             Inches(0.9),
             "Quando o reuso aparece, a próxima conversa em comitê "
             "muda. Não é “vamos criar mais um agente”. É “qual "
             "capacidade está faltando?”",
             font=FONT_BODY, size=18, color=MS_DARK, italic=True)


def slide_demo4(slide):
    slide_demo(
        slide,
        num=4,
        title="Foundry portal — segundo agente, mesmas tools",
        bullets=[
            "Prompt agent hr-orchestrator consumindo as mesmas seis tools MCP.",
            "Turno: “quero solicitar 18 dias de férias em julho”.",
            "Mesma sequência de tool calls. Resultado diferente porque a "
            "política e a permissão são diferentes — as skills são iguais.",
        ],
        return_cue=(
            "A última pergunta é a que mais incomoda em comitê."
        ),
    )


def slide_antipadroes(slide):
    add_title(slide, "Anti-padrões — o que eu vejo dar errado")
    items = [
        ("vanity agent",
         "Cada diretoria pede “o meu agente”. Quando você abre, é a mesma "
         "busca + a mesma política, com outro nome. Mostre o agente vizinho."),
        ("skill-shaped agent",
         "Chamam de “agente” o que é só uma função tipada. Promova para "
         "skill, registre, versione. Vira ativo da empresa."),
        ("orchestrator-shaped skill",
         "Skill que internamente faz três coisas (chama API, formata, "
         "escreve em banco). Quebre em três. Cada uma tem dono diferente."),
        ("silent retry",
         "Skill que engole erro e devolve OK “vazio”. Use "
         "TEMPORARY_FAILURE. O agente decide se tenta de novo."),
    ]
    y = Inches(2.2)
    for chip_label, body in items:
        chip, chip_w = add_chip(slide, MARGIN_X, y, chip_label,
                                fill=MS_WARN, color=WHITE, size=11)
        add_text(slide, MARGIN_X + chip_w + Inches(0.15),
                 y - Inches(0.03), CONTENT_W - chip_w - Inches(0.3),
                 Inches(1.0), body,
                 font=FONT_BODY, size=15, color=MS_INK)
        y = y + Inches(1.1)


def slide_pergunta_central(slide):
    add_text(slide, MARGIN_X, Inches(1.4), CONTENT_W, Inches(0.5),
             "A pergunta que define o roadmap dos próximos 12 meses",
             font=FONT_HEAD, size=20, color=MS_MUTED,
             align=PP_ALIGN.CENTER)

    add_paragraphs(
        slide, MARGIN_X, Inches(2.4), CONTENT_W, Inches(3.4),
        [
            ("A empresa precisa de",
             {"size": 40, "color": MS_DARK, "align": PP_ALIGN.CENTER}),
            ("mais agentes",
             {"size": 64, "bold": True, "color": MS_INK,
              "align": PP_ALIGN.CENTER, "space_before": 6}),
            ("ou de",
             {"size": 40, "color": MS_DARK, "align": PP_ALIGN.CENTER,
              "space_before": 6}),
            ("melhores capacidades reutilizáveis?",
             {"size": 64, "bold": True, "color": MS_BLUE,
              "align": PP_ALIGN.CENTER, "space_before": 6}),
        ],
        line_spacing=1.1,
    )

    add_text(slide, MARGIN_X, Inches(6.2), CONTENT_W, Inches(0.5),
             "A resposta honesta para a maioria das empresas hoje é a "
             "segunda — e ninguém está medindo isso.",
             font=FONT_BODY, size=16, color=MS_MUTED, italic=True,
             align=PP_ALIGN.CENTER)


def slide_segunda_feira(slide):
    add_title(slide, "O que levar para segunda-feira")
    add_paragraphs(
        slide, MARGIN_X, BODY_TOP, CONTENT_W, BODY_H,
        [
            ("1.  Inventário de agentes. Liste o que já existe. Pinte os duplicados.",
             {"size": 17, "color": MS_INK}),
            ("2.  Inventário de capacidades. Quais skills os agentes existentes chamam?",
             {"size": 17, "color": MS_INK, "space_before": 8}),
            ("3.  Registry mínimo. Comece com 5–6 capacidades verdadeiras. Não 50.",
             {"size": 17, "color": MS_INK, "space_before": 8}),
            ("4.  Contrato escrito. SkillManifest + ResultCategory em um README.",
             {"size": 17, "color": MS_INK, "space_before": 8}),
            ("5.  Primeiro teste de governança. “Sem owner, não registra.” Vire regra.",
             {"size": 17, "color": MS_INK, "space_before": 8}),
            ("6.  Foundry como destino, não como ponto de partida. Prompt agents "
             "em cima do MCP server, não o contrário.",
             {"size": 17, "color": MS_INK, "space_before": 8}),
        ],
        line_spacing=1.35,
    )
    quote_box = _solid_rect(slide, MARGIN_X, Inches(6.05),
                            CONTENT_W, Inches(0.85), MS_GRAY)
    quote_box.line.fill.background()
    _solid_rect(slide, MARGIN_X, Inches(6.05), Inches(0.07),
                Inches(0.85), MS_BLUE)
    add_text(slide, Inches(0.95), Inches(6.18), CONTENT_W - Inches(0.5),
             Inches(0.7),
             "O próximo agente que vocês forem aprovar vale a pergunta: "
             "“qual skill nova esse agente exige?” Se a resposta for "
             "“nenhuma”, talvez ele já exista.",
             font=FONT_BODY, size=15, color=MS_DARK, italic=True)


def slide_closing(slide):
    add_paragraphs(
        slide, MARGIN_X, Inches(2.4), CONTENT_W, Inches(3.0),
        [
            ("Menos agentes.",
             {"size": 84, "bold": True, "color": WHITE,
              "align": PP_ALIGN.CENTER}),
            ("Mais capacidades.",
             {"size": 84, "bold": True, "color": MS_ACCENT,
              "align": PP_ALIGN.CENTER, "space_before": 8}),
        ],
        line_spacing=1.05,
    )
    _solid_rect(slide, Inches(5.7), Inches(5.4), Inches(2.0),
                Inches(0.06), MS_ACCENT)
    add_paragraphs(
        slide, MARGIN_X, Inches(5.7), CONTENT_W, Inches(1.4),
        [
            ("Ricardo Cataldi · Microsoft GBB",
             {"size": 18, "color": WHITE, "align": PP_ALIGN.CENTER,
              "bold": True}),
            (REPO_URL,
             {"size": 16, "color": MS_ACCENT, "italic": True,
              "align": PP_ALIGN.CENTER, "space_before": 6}),
            ("AI Tech Circle · 28 de maio de 2026",
             {"size": 14, "color": MS_GRAY,
              "align": PP_ALIGN.CENTER, "space_before": 6}),
        ],
        line_spacing=1.2,
    )


# ----------------------------------------------------------------------
# Deck assembly
# ----------------------------------------------------------------------


# (builder_fn, style)  — style ∈ {"cover","light","dark","question","closing"}
DECK = [
    (build_cover, "cover"),
    (slide_pergunta_honesta, "light"),
    (slide_armadilha, "light"),
    (slide_vocabulario, "light"),
    (slide_decomposicao, "light"),
    (slide_demo1, "dark"),
    (slide_contrato, "light"),
    (slide_envelope, "light"),
    (slide_demo2, "dark"),
    (slide_governanca, "light"),
    (slide_demo3, "dark"),
    (slide_reuso, "light"),
    (slide_demo4, "dark"),
    (slide_antipadroes, "light"),
    (slide_pergunta_central, "question"),
    (slide_segunda_feira, "light"),
    (slide_closing, "closing"),
]


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    blank = prs.slide_layouts[6]
    total = len(DECK)

    for i, (builder, style) in enumerate(DECK, start=1):
        slide = prs.slides.add_slide(blank)

        if style == "cover":
            builder(slide, total)
            continue

        if style == "dark":
            dark_slide(slide, total, i)
        elif style == "question":
            question_slide(slide, total, i)
        elif style == "closing":
            _paint_background(slide, MS_DEEP)
            add_header_footer(slide, i, total, dark=True)
        else:
            light_slide(slide, total, i)

        # Style-specific builders take only the slide
        if builder.__code__.co_argcount == 1:
            builder(slide)
        else:
            builder(slide, total)

    out = Path(__file__).parent / "skill-first-ai.pptx"
    prs.save(out)
    return out


if __name__ == "__main__":
    path = build()
    print(f"wrote {path}")
